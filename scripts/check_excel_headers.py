import openpyxl, os, json

files = [
    (r"D:\Files\014数据文件夹\13-【井书·独家资料包】河南\河南全套\24理科录取统计Excel版本\24河南理科录取统计所有批次.xlsx", "24理科录取统计所有批次"),
    (r"D:\Files\014数据文件夹\13-【井书·独家资料包】河南\河南全套\24文科录取统计Excel版本\文科所有批次.xlsx", "24文科录取统计所有批次"),
    (r"D:\Files\014数据文件夹\2026河南高考志愿填报_参考文档\招生计划\H河南-招生计划-2024.xlsx", "H河南-招生计划-2024"),
    (r"D:\Files\014数据文件夹\2026河南高考志愿填报_参考文档\招生计划\20250621-河南-2025-招生计划.xlsx", "20250621-河南-2025-招生计划"),
    (r"D:\Files\014数据文件夹\G25.河南——98数据\河南-2023-招生计划.xlsx", "河南-2023-招生计划"),
    (r"D:\Files\014数据文件夹\2026河南高考志愿填报_参考文档\录取统计\河南省-2024年专业录取分数.xlsx", "河南省-2024年专业录取分数"),
]

for fpath, fname in files:
    print(f"\n=== {fname} ===")
    if not os.path.exists(fpath):
        print("  FILE NOT FOUND")
        continue
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        print(f"  Sheets: {wb.sheetnames}")
        for sn in wb.sheetnames[:3]:
            ws = wb[sn]
            print(f"  [{sn}] max_row: {ws.max_row}, max_col: {ws.max_column}")
            # Get header row (first row)
            headers = []
            for col in range(1, min(ws.max_column or 100, 50) + 1):
                cell = ws.cell(1, col)
                headers.append(str(cell.value) if cell.value is not None else "")
            print(f"  Headers: {headers}")
            # Get sample row 2
            if ws.max_row and ws.max_row >= 2:
                row2 = []
                for col in range(1, min(ws.max_column or 100, 50) + 1):
                    cell = ws.cell(2, col)
                    row2.append(str(cell.value) if cell.value is not None else "")
                print(f"  Row 2: {row2}")
        wb.close()
    except Exception as e:
        print(f"  ERROR: {e}")
