"""
提取学科评估和大学排名数据，保存为 JSON
"""
import json, os, re
import openpyxl

OUT_DIR = r'D:\GoakaoProject\docs'
os.makedirs(OUT_DIR, exist_ok=True)

# ========== Part 1: 学科评估 ==========
print('=' * 60)
print('Part 1: 提取第四轮学科评估数据')
print('=' * 60)

subject_fp = r'D:\Files\014数据文件夹\2315\GK资料大全\第四轮全国高校学科评估结果Excel表格.xlsx'
wb = openpyxl.load_workbook(subject_fp, data_only=True)
ws = wb.active

# 解析 header (row 2)
headers = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
print(f'Headers: {headers}')

records = []
for row in ws.iter_rows(min_row=3, values_only=True):
    subject_code_name = str(row[1]).strip() if row[1] else ''
    level = str(row[2]).strip() if row[2] else ''
    school_raw = str(row[3]).strip() if row[3] else ''
    category = str(row[4]).strip() if row[4] else ''
    major_class = str(row[5]).strip() if row[5] else ''
    
    if not subject_code_name or not level or not school_raw:
        continue
    
    # 解析学科代码和名称: "0101哲学" → code=0101, name=哲学
    subject_match = re.match(r'^(\d+)(.+)$', subject_code_name)
    subject_code = subject_match.group(1) if subject_match else ''
    subject_name = subject_match.group(2) if subject_match else subject_code_name
    
    # 解析学校代码和名称: "10001北京大学_x000D_\n" → code=10001, name=北京大学
    school_clean = re.sub(r'_x000D_\s*', '', school_raw).strip()
    school_clean = re.sub(r'\s+', '', school_clean)
    school_match = re.match(r'^(\d+)(.+)$', school_clean)
    school_code = school_match.group(1) if school_match else ''
    school_name = school_match.group(2) if school_match else school_clean
    
    records.append({
        'school': school_name,
        'school_code': school_code,
        'subject': subject_name,
        'subject_code': subject_code,
        'level': level,
        'category': category,
        'major_class': major_class,
        'year': 2017,
    })

print(f'  Total records: {len(records)}')
print(f'  Schools: {len(set(r["school"] for r in records))}')
print(f'  Subjects: {len(set(r["subject"] for r in records))}')
print(f'  Levels: {sorted(set(r["level"] for r in records))}')

# 统计各等级的学校-学科对
from collections import Counter
level_counts = Counter(r['level'] for r in records)
for lv, cnt in sorted(level_counts.items()):
    print(f'    {lv}: {cnt}')

output = {
    'meta': {
        'version': '1.0',
        'source': '第四轮全国高校学科评估结果Excel表格.xlsx',
        'total_records': len(records),
        'year': 2017,
        'description': '第四轮全国高校学科评估结果（2017年公布）',
    },
    'data': records,
}

out_path = os.path.join(OUT_DIR, 'subject_evaluation.json')
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False)
print(f'  Saved to: {out_path}')
print(f'  File size: {os.path.getsize(out_path)/1024:.1f} KB')

# ========== Part 2: 大学排名 ==========
print()
print('=' * 60)
print('Part 2: 提取大学排名数据（软科）')
print('=' * 60)

ranking_records = []

# --- 2023 软科 ---
print('\n--- 2023 软科 ---')
fp_2023 = r'D:\Files\014数据文件夹\河南--2024年志愿填报资料包（最新）\全国通用高考数据\1、大学排名\2023软科排名总榜.xlsx'
wb3 = openpyxl.load_workbook(fp_2023, data_only=True)
ws3 = wb3.active
for row in ws3.iter_rows(min_row=2, values_only=True):
    if row[1] is None:
        continue
    school = str(row[1]).strip()
    rank_str = str(row[2]).strip() if row[2] else ''
    try:
        rank = int(rank_str)
    except ValueError:
        continue
    tags = str(row[3]).strip() if row[3] else ''
    school_type = str(row[4]).strip() if row[4] else ''
    region = str(row[5]).strip() if row[5] else ''
    score_str = str(row[6]).strip() if row[6] else ''
    try:
        score = float(score_str)
    except ValueError:
        score = None
    
    ranking_records.append({
        'school': school,
        'ranking': rank,
        'score': score,
        'year': 2023,
        'source': '软科',
        'region': region,
        'type': school_type,
        'tags': tags,
    })
print(f'  2023 records: {len([r for r in ranking_records if r["year"]==2023])}')

# --- 2022 软科 ---
print('\n--- 2022 软科 ---')
fp_2022 = r'D:\Files\014数据文件夹\河南--2024年志愿填报资料包（最新）\全国通用高考数据\1、大学排名\2022软科排名总榜.xlsx'
wb2 = openpyxl.load_workbook(fp_2022, data_only=True)
# 使用主榜
ws2 = wb2['2022 主榜']
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row[1] is None:
        continue
    rank = row[0]
    school = str(row[1]).strip()
    region = str(row[2]).strip() if row[2] else ''
    school_type = str(row[3]).strip() if row[3] else ''
    score = row[4] if row[4] else None
    
    ranking_records.append({
        'school': school,
        'ranking': int(rank) if rank else None,
        'score': float(score) if score else None,
        'year': 2022,
        'source': '软科',
        'region': region,
        'type': school_type,
        'tags': '',
    })
print(f'  2022 records: {len([r for r in ranking_records if r["year"]==2022])}')

# --- 2021 软科 ---
print('\n--- 2021 软科 ---')
fp_2021 = r'D:\Files\014数据文件夹\河南--2024年志愿填报资料包（最新）\全国通用高考数据\1、大学排名\2021软科排名总榜.xlsx'
wb1 = openpyxl.load_workbook(fp_2021, data_only=True)
ws1 = wb1.active
for row in ws1.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    school = str(row[0]).strip()
    rank = row[2]
    region = str(row[3]).strip() if row[3] else ''
    school_type = str(row[5]).strip() if row[5] else ''
    score_raw = row[7]
    try:
        score = float(score_raw) if score_raw and str(score_raw).strip() != '-' else None
    except (ValueError, TypeError):
        score = None
    
    ranking_records.append({
        'school': school,
        'ranking': int(rank) if rank else None,
        'score': score,
        'year': 2021,
        'source': '软科',
        'region': region,
        'type': school_type,
        'tags': '',
    })
print(f'  2021 records: {len([r for r in ranking_records if r["year"]==2021])}')

# --- 2020 大学排名 ---
print('\n--- 2020 大学排名 ---')
fp_2020 = r'D:\Files\014数据文件夹\G25.河南——98数据\河南-历年高考数据\2020年大学排名.xlsx'
try:
    wb0 = openpyxl.load_workbook(fp_2020, data_only=True)
    ws0 = wb0.active
    print(f'  Sheets: {wb0.sheetnames}, Rows: {ws0.max_row}, Cols: {ws0.max_column}')
    for row in ws0.iter_rows(min_row=1, max_row=3, values_only=True):
        print(f'  Sample: {list(row)}')
    
    # Detect header and parse
    for i, row in enumerate(ws0.iter_rows(min_row=1, max_row=5, values_only=True)):
        vals = [str(v) if v else '' for v in row]
        if any('排名' in v for v in vals) and any('大学' in v or '学校' in v for v in vals):
            header_row = i + 1
            print(f'  Header at row {header_row}: {vals}')
            # Parse from next row
            for data_row in ws0.iter_rows(min_row=header_row+1, values_only=True):
                if data_row[0] is None:
                    continue
                # Try to find rank and school columns
                rank_val = None
                school_val = None
                score_val = None
                for j, v in enumerate(data_row):
                    vs = str(v).strip() if v else ''
                    if vs.isdigit() and rank_val is None and j < 2:
                        rank_val = int(vs)
                    elif '大学' in vs or '学院' in vs:
                        school_val = vs
                    elif vs.replace('.', '').isdigit() and score_val is None:
                        try:
                            score_val = float(vs)
                        except:
                            pass
                if school_val:
                    ranking_records.append({
                        'school': school_val,
                        'ranking': rank_val,
                        'score': score_val,
                        'year': 2020,
                        'source': '通用排名',
                        'region': '',
                        'type': '',
                        'tags': '',
                    })
            break
    print(f'  2020 records: {len([r for r in ranking_records if r["year"]==2020])}')
except Exception as e:
    print(f'  2020 file error: {e}')

print(f'\nTotal ranking records: {len(ranking_records)}')
print(f'Years: {sorted(set(r["year"] for r in ranking_records))}')
print(f'Schools covered: {len(set(r["school"] for r in ranking_records))}')

ranking_output = {
    'meta': {
        'version': '1.0',
        'source': '软科排名（2021-2023）+ 通用大学排名（2020）',
        'total_records': len(ranking_records),
        'years': sorted(set(r['year'] for r in ranking_records)),
    },
    'data': ranking_records,
}

out_path2 = os.path.join(OUT_DIR, 'university_rankings.json')
with open(out_path2, 'w', encoding='utf-8') as f:
    json.dump(ranking_output, f, ensure_ascii=False)
print(f'Saved to: {out_path2}')
print(f'File size: {os.path.getsize(out_path2)/1024:.1f} KB')

print()
print('=' * 60)
print('提取完成！')
print('=' * 60)
